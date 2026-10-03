import requests
import json
import re
import os
from docx import Document

# ================= CONFIGURATION =================
API_KEY = os.getenv("OPENROUTER_API_KEY") 
MODEL = "deepseek/deepseek-v4.1-flash" 
BASE_URL = "https://openrouter.ai/api/v1/chat/completions"
OUTPUT_FOLDER = "procedures"

if not API_KEY:
    raise EnvironmentError(
        "ERROR: OPENROUTER_API_KEY environment variable not found."
    )

SYSTEM_PROMPT = """Role: Act as an expert Fintech Systems Architect and Head of Payment Operations with deep expertise in global payment rails (ISO 20022, SWIFT, ACH, RTP) and banking regulatory compliance.
Task: Develop a comprehensive set of 30 Standard Operating Procedures (SOPs) for a fictional high-scale payment processing ecosystem. You will generate these procedures in batches of 5.

Fictional System Architecture:
OmniFlow: API Gateway, ingestion, and webhook delivery.
AegisPay Core: Account registry, master data, and validation engine.
NexusRoute: Network routing, operator selection, and transformation.
SentinelTrace: Fraud monitoring, anomaly detection, and real-time tracking.
MeridianClear: Settlement preparation, sequencing, and execution.
CobaltSettlement: Reconciliation, variance matching, and exception queues.
VaultGuard: Liquidity management, FX rate sourcing, and funding pools.
IronClad: Immutable audit archiving, compliance gateway, and forensic logging.

Required Template for Every Procedure:
Each SOP must be formatted with these exact headers:
Procedure ID & Title: (e.g., PAY-OPS-XXX: [Title])
Purpose & Scope: ...
Applicable Payment Corridors & Networks: ...
Invented Processing Systems & Architecture: ...
Detailed Key Steps: ...
Bank Account Number Handling: ...
Special Currency Handling: ...
Escalation Framework: ...
RACI Matrix & Responsible Roles: ...
Compliance, Regulatory & Audit References: ...
System Integration, Data Flow & API/Webhook Touchpoints: ...
SLA, Cut-off Times & Settlement Windows: ...
Risk Controls, Exception Handling & Reconciliation: ...
Version Control, Approval Workflow & Change Management: ...

The 30 Subject Headers to be Covered:
1. Domestic ACH Origination & Validation
2. SWIFT MT/MX Message Processing & ISO 20022 Field Mapping
3. Real-Time Payments (RTP) & Instant Transfer Handling
4. Batch Processing & End-of-Day Settlement
5. Multi-Currency Account Reconciliation & Mapping
6. FX Settlement, Conversion & Liquidity Allocation
7. Payment Dispute, Return & Chargeback Workflows
8. Fraud Detection, Transaction Monitoring & Exception Handling
9. Intraday Liquidity Management & Funding Optimization
10. Regulatory Reporting (Transaction, Tax, Sanctions Screening)
11. Third-Party, Payroll & Vendor Payment Processing
12. Payment Cut-off Management & Holiday Calendar Handling
13. Audit Readiness, Exception Logging & Control Testing
14. Business Continuity, DR Failover & Payment System Redundancy
15. Merchant Acquiring & Payout Network Settlement
16. Cross-Border Correspondent Banking & Nostro/Vostro Reconciliation
17. Digital Wallet & Tokenized Payment Integration
18. API Payment Flows & Webhook Reconciliation
19. Regulatory Sandbox Testing & Pre-Production Validation
20. Advanced Account Validation & IBAN/BIC Normalization Engine
21. High-Value Wire (Fedwire/CHIPS) Origination & Settlement
22. Cross-Border FX & Capital Control Compliance
23. ACH Returns, Reversals & Recurring Payment Management
24. Treasury Management & Cash Pooling Operations
25. Payment Fraud & Cyber Incident Response Protocol
26. SWIFT gpi Tracking & End-to-End Payment Monitoring
27. Intraday Credit & Liquidity Optimization Strategies
28. Payment Network Switching & Failover Routing
29. Client Onboarding & KYC/AML Integration for Payments
30. Payment Data Governance & Privacy Compliance

Execution Instructions:
Use highly formal, technical language (e.g., "mod-10 checksum," "pacs.008," "idempotency keys")."""

# ================= HELPER FUNCTIONS =================

def call_openrouter(batch_start, batch_end):
    # Explicitly tell the AI how to separate the documents so Python can split them
    prompt = (f"Generate procedures {batch_start} through {batch_end}. "
              f"CRITICAL: You must start each procedure with 'Procedure ID & Title: PAY-OPS-XXX'. "
              f"Do not wrap the ID in markdown bolding or hashtags. Ensure there is a clear line break between procedures.")
    
    headers = {
        "Authorization": f"Bearer {API_KEY}",
        "Content-Type": "application/json",
        "HTTP-Referer": "http://localhost:3000", 
        "X-Title": "Fintech SOP Generator",
    }
    
    payload = {
        "model": MODEL,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": prompt}
        ],
        "temperature": 0
    }

    response = requests.post(BASE_URL, headers=headers, data=json.dumps(payload))
    if response.status_code == 200:
        return response.json()['choices'][0]['message']['content']
    else:
        print(f"Server Response: {response.text}")
        raise Exception(f"API Error: {response.status_code}")

def save_to_word(text, filename):
    filepath = os.path.join(OUTPUT_FOLDER, filename)
    doc = Document()
    # Clean up the text slightly (remove markdown artifacts like **)
    clean_text = text.replace('**', '') 
    doc.add_paragraph(clean_text)
    doc.save(filepath)

def append_to_master(text, master_doc_obj):
    # Remove markdown bolding for the master doc as well
    clean_text = text.replace('**', '')
    master_doc_obj.add_paragraph(clean_text)
    master_doc_obj.add_page_break()

# ================= MAIN EXECUTION =================

def main():
    if not os.path.exists(OUTPUT_FOLDER):
        os.makedirs(OUTPUT_FOLDER)
        print(f"Created folder: {OUTPUT_FOLDER}")

    master_doc = Document()
    master_doc.add_heading('Comprehensive Payment Operations SOP Manual', 0)
    
    total_procedures = 30
    batch_size = 5
    
    for i in range(0, total_procedures, batch_size):
        start = i + 1
        end = i + batch_size
        
        print(f"\n--- Processing Batch: Procedures {start} to {end} ---")
        try:
            raw_content = call_openrouter(start, end)
            
            sops = re.split(r'(?=\s*(?:#*\s*Procedure ID & Title:\s*PAY-OPS-))', raw_content)
            sops = [sop.strip() for sop in sops if sop.strip()]

            print(f"Detected {len(sops)} procedures in this batch.")

            for sop_text in sops:
                match = re.search(r'(PAY-OPS-\d+)', sop_text)
                if match:
                    doc_id = match.group(1)
                    filename = f"{doc_id}.docx"
                    save_to_word(sop_text, filename)
                    print(f"Successfully saved individual file: {filename}")
                else:
                    print("Warning: Could not find a PAY-OPS ID in this section. Skipping single file save.")
                
                append_to_master(sop_text, master_doc)

            if end < total_procedures:
                user_input = input("\nBatch complete. Type 'Next' to proceed: ")
                if user_input.lower() != "next":
                    break

        except Exception as e:
            print(f"An error occurred: {e}")
            break

    master_filepath = "Master_Procedure.docx"
    master_doc.save(master_filepath)
    print(f"\nProcess finished.")

if __name__ == "__main__":
    main()
