"""
FILE: OperationalRAG_Prototype.py
PURPOSE: DEMO project:
        (1) Enter the full path to your procedures folder, 
        (2) Enter a natural language question 
        (3) The system will provide an answer and ask if you wish to open the cited .docx files immediately for manual verification.
        (4) Every query is automatically saved as a timestamped document in the /Log folder.
INPUTS: ./procedures/{30 x Procedures.docx}, OPENROUTER_API_KEY set in system environment variables
OUTPUTS: ./Log/Query_%Y%m%d%H%M%S.docx
DEPENDENCIES: numpy, python-docx, sentence-transformers, openai
"""

import os
import glob
import re
import numpy as np
import getpass
from datetime import datetime 
from docx import Document
from sentence_transformers import SentenceTransformer
from openai import OpenAI

# ==========================================
# RAG ENGINE CLASS (Adapted from RunRag.py)
# ==========================================
class OperationalRAG:
    def __init__(self, folder_path):
        self.folder_path = folder_path
        self.embedder = SentenceTransformer('all-mpnet-base-v2')
        self.client = self._get_client()
        
        # Initialize Knowledge Base
        self.raw_docs = self._load_docs()
        self.chunks_data = self._chunk_docs(self.raw_docs)
        self.chunks_text = [c["text"] for c in self.chunks_data]
        self.vector_matrix = self.embedder.encode(self.chunks_text, normalize_embeddings=True)

    def _get_client(self):
        key = os.environ.get('OPENROUTER_API_KEY') or getpass.getpass('Paste your OpenRouter API key: ')
        return OpenAI(base_url='https://openrouter.ai/api/v1', api_key=key)

    def _load_docs(self):
        docs = []
        files = glob.glob(os.path.join(self.folder_path, "*.docx"))
        if not files:
            raise FileNotFoundError(f"No .docx files found in {self.folder_path}")
        for filepath in files:
            filename = os.path.basename(filepath)
            doc = Document(filepath)
            text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
            if text.strip():
                docs.append((text, filename)) 
        return docs

    def _chunk_docs(self, docs, size=250, overlap=50):
        all_chunks = []
        for text, source in docs:
            words = text.split()
            i = 0
            while i < len(words):
                chunk_text = ' '.join(words[i:i+size])
                all_chunks.append({"text": chunk_text, "source": source})
                if i + size >= len(words): break
                i += max(1, size - overlap)
        return all_chunks

    def retrieve(self, question, k=15):
        q_vec = self.embedder.encode([question], normalize_embeddings=True)[0]
        scores = self.vector_matrix @ q_vec
        top_k_idx = np.argsort(-scores)[:k]
        return [(self.chunks_data[i]["text"], self.chunks_data[i]["source"]) for i in top_k_idx]

    def generate(self, question, context):
        context_text = "\n\n".join(f"Source {src}: {txt}" for txt, src in context)
        system_prompt = (
            "Answer using ONLY the notes provided. "
            "You MUST explicitly state BOTH the Document Name and the Procedure Reference (e.g., PAY-OPS-XXX) you used for your answer. "
            "Answer based on the provided notes. If the information is partially available, provide what is there and cite the source. Only if the information is completely absent should you say 'The notes do not say."
            "Do not use any other knowledge."
        )
        response = self.client.chat.completions.create(
            model="openai/gpt-4o-mini",
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": f"NOTES:\n{context_text}\n\nQUESTION: {question}"}
            ],
            temperature=0.0
        )
        return response.choices[0].message.content.strip()

# ==========================================
# UTILITY FUNCTIONS FOR LOGGING & UX
# ==========================================
def save_to_audit_log(question, answer, sources):
    """Saves the RAG interaction to a Word document with timestamped filename."""
    if not os.path.exists("Log"):
        os.makedirs("Log")
    
    doc = Document()
    doc.add_heading('OperationalRAG Audit Log', 0)
    
    doc.add_heading('User Question:', level=1)
    doc.add_paragraph(question)
    
    doc.add_heading('AI Generated Answer:', level=1)
    doc.add_paragraph(answer)
    
    doc.add_heading('Supporting Evidence (Retrieved Chunks):', level=1)
    for src, text in sources:
        doc.add_paragraph(f"File: {src}\nContent: {text}\n{'-'*30}")
    
    doc.add_paragraph("\n\nDISCLAIMER: This is an AI-generated response. "
                      "The user MUST manually verify this information against the original "
                      "procedural documents in the source folder before taking operational action.")
    
    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    filename = f"Log/Query_{timestamp}.docx"
    doc.save(filename)
    return filename

def open_file(filepath):
    """Opens the file using the system default application."""
    try:
        os.startfile(filepath) # Windows
    except AttributeError:
        import subprocess
        opener = "open" if os.uname().sysname == "Darwin" else "xdg-open"
        subprocess.call([opener, filepath])

# ==========================================
# MAIN APPLICATION LOOP
# ==========================================
def main():
    print("--- Welcome to OperationalRAG System ---")
    folder = input("1) Enter the path to the folder containing Procedures (.docx): ").strip()
    
    try:
        rag_system = OperationalRAG(folder)
        print("\n[System] Knowledge base indexed successfully.")
    except Exception as e:
        print(f"Error initializing system: {e}")
        return

    while True:
        query = input("\n2) Enter your question (or type 'exit' to quit): ").strip()
        if query.lower() == 'exit': break
        if not query: continue

        # RAG Process
        context_chunks = rag_system.retrieve(query) 
        answer = rag_system.generate(query, context_chunks)
        
        # Display results to user
        print("\n" + "="*50)
        print(f"ANSWER: {answer}")
        print("="*50)
        print("REMINDER: Please verify this information in the procedures folder. AI may be incorrect.")

        # Audit Logging (background)
        log_path = save_to_audit_log(query, answer, context_chunks)
        print(f"\n[Audit] Response logged to: {log_path}")

        # --- HANDLE CITED DOCUMENT VERIFICATION ---
        if context_chunks:
            # 1. Extract all Procedure IDs (e.g., PAY-OPS-001) mentioned in the generated answer
            cited_ids = re.findall(r'[A-Z]{3}-[A-Z]{3}-\d+', answer)
            
            # 2. Map these IDs to actual filenames from our retrieved context
            docs_to_open = []
            unique_retrieved_files = list(set([src for text, src in context_chunks]))
            
            for file_name in unique_retrieved_files:
                # If the filename contains any of the IDs cited in the answer, add it to the open list
                if any(cite_id in file_name for cite_id in cited_ids):
                    docs_to_open.append(file_name)
            
            if docs_to_open:
                print("\nProcedures cited in the answer:")
                for i, doc_name in enumerate(docs_to_open, 1):
                    print(f" {i}. {doc_name}")
                
                # 3. Strict loop for Y/N prompt
                while True:
                    choice = input(f"\nWould you like to open all {len(docs_to_open)} cited procedure(s) for verification? (Y/N): ").strip().upper()
                    if choice == 'Y':
                        for doc_name in docs_to_open:
                            full_path = os.path.join(folder, doc_name)
                            open_file(full_path)
                        break # Exit loop after opening files
                    elif choice == 'N':
                        print("Skipping file verification.")
                        break # Exit loop normally
                    else:
                        print("Invalid input. Please enter 'Y' for Yes or 'N' for No.")

    print("\nSystem closed. Goodbye!")

if __name__ == "__main__":
    main()
